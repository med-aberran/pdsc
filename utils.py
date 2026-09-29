# -*- coding: utf-8 -*-
"""
utils.py
========
Fonctions utilitaires reutilisees par routes.py : upload de fichiers,
generation de QR codes, generation de PDF et d'Excel, envoi d'e-mails,
jetons securises, journalisation et decorateurs de controle d'acces par role.
"""

import io
import os
import random
import string
import uuid
from datetime import datetime
from functools import wraps

import mimetypes

import qrcode
import requests
from flask import abort, current_app, flash, redirect, request, session, url_for
from flask_login import current_user
from flask_mail import Message
from itsdangerous import URLSafeTimedSerializer
from marshmallow import Schema, fields
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from werkzeug.utils import secure_filename


# ---------------------------------------------------------------------------
# Schemas Marshmallow (serialisation pour l'API JWT)
# ---------------------------------------------------------------------------

class ServiceRequestSchema(Schema):
    """Serialise une ServiceRequest pour les reponses JSON de l'API."""
    tracking_number = fields.Str()
    service = fields.Method('get_service_name')
    status = fields.Str()
    status_label = fields.Method('get_status_label')
    created_at = fields.DateTime()
    updated_at = fields.DateTime()

    def get_service_name(self, obj):
        return obj.service.name_fr

    def get_status_label(self, obj):
        return obj.status_label()


request_schema = ServiceRequestSchema()
requests_schema = ServiceRequestSchema(many=True)


# ---------------------------------------------------------------------------
# Internationalisation
# ---------------------------------------------------------------------------

def get_locale():
    """Determine la langue active : parametre URL > session > preference du
    compte > navigateur > langue par defaut."""
    requested = request.args.get('lang')
    if requested in current_app.config['LANGUAGES']:
        session['lang'] = requested
        if current_user.is_authenticated:
            current_user.preferred_language = requested
        return requested
    if session.get('lang') in current_app.config['LANGUAGES']:
        return session['lang']
    if current_user.is_authenticated and current_user.preferred_language:
        return current_user.preferred_language
    return request.accept_languages.best_match(current_app.config['LANGUAGES']) \
        or current_app.config['BABEL_DEFAULT_LOCALE']


# ---------------------------------------------------------------------------
# Upload de fichiers (Supabase Storage en production, disque en local)
# ---------------------------------------------------------------------------
# En base, on continue de stocker un chemin relatif du type
# "news/3f2a...png" : le code métier et les colonnes ne changent pas.
# Les fichiers "requests/..." vont dans le bucket PRIVE, le reste dans le
# bucket PUBLIC.

PRIVATE_PREFIXES = ('requests/',)
STORAGE_TIMEOUT = 20


def allowed_file(filename, allowed_extensions):
    return '.' in filename and \
        filename.rsplit('.', 1)[1].lower() in allowed_extensions


def storage_enabled():
    """True si Supabase Storage est configure (URL + cle service_role)."""
    cfg = current_app.config
    return bool(cfg.get('SUPABASE_URL') and cfg.get('SUPABASE_KEY'))


def _is_private(relative_path):
    return relative_path.startswith(PRIVATE_PREFIXES)


def _bucket_for(relative_path):
    cfg = current_app.config
    return cfg['SUPABASE_PRIVATE_BUCKET'] if _is_private(relative_path) \
        else cfg['SUPABASE_BUCKET']


def _storage_headers(content_type=None):
    key = current_app.config['SUPABASE_KEY']
    headers = {'Authorization': f'Bearer {key}', 'apikey': key}
    if content_type:
        headers['Content-Type'] = content_type
    return headers


def _require_storage_on_vercel():
    if os.environ.get('VERCEL') and not storage_enabled():
        raise RuntimeError(
            "Supabase Storage n'est pas configure : definir SUPABASE_URL et "
            "SUPABASE_KEY dans les variables d'environnement Vercel "
            "(le disque de Vercel est en lecture seule)."
        )


def _storage_put(relative_path, data, content_type):
    """Envoie des octets vers Supabase Storage. Leve une exception en cas d'echec."""
    cfg = current_app.config
    url = f"{cfg['SUPABASE_URL']}/storage/v1/object/{_bucket_for(relative_path)}/{relative_path}"
    resp = requests.post(
        url, data=data, headers=_storage_headers(content_type), timeout=STORAGE_TIMEOUT)
    if resp.status_code >= 300:
        current_app.logger.error('Upload Supabase refuse (%s) : %s', resp.status_code, resp.text)
        raise RuntimeError('Echec de l\'envoi du fichier vers le stockage.')


def _storage_delete(relative_path):
    cfg = current_app.config
    url = f"{cfg['SUPABASE_URL']}/storage/v1/object/{_bucket_for(relative_path)}/{relative_path}"
    try:
        resp = requests.delete(url, headers=_storage_headers(), timeout=STORAGE_TIMEOUT)
        if resp.status_code >= 300:
            current_app.logger.warning(
                'Suppression Supabase impossible (%s) : %s', resp.status_code, resp.text)
    except requests.RequestException as exc:
        current_app.logger.warning('Suppression Supabase impossible : %s', exc)


def save_bytes(data, relative_path, content_type='application/octet-stream'):
    """Enregistre des octets (Supabase Storage ou disque local)."""
    _require_storage_on_vercel()
    if storage_enabled():
        _storage_put(relative_path, data, content_type)
    else:
        full_path = os.path.join(current_app.config['UPLOAD_FOLDER'], relative_path)
        os.makedirs(os.path.dirname(full_path), exist_ok=True)
        with open(full_path, 'wb') as fh:
            fh.write(data)
    return relative_path


def save_uploaded_file(file_storage, subfolder=''):
    """Sauvegarde un fichier uploade avec un nom unique et securise.
    Retourne le chemin relatif (ex. 'news/abc.png') ou None."""
    if not file_storage or not file_storage.filename:
        return None

    original_name = secure_filename(file_storage.filename)
    ext = original_name.rsplit('.', 1)[1].lower() if '.' in original_name else ''
    unique_name = f'{uuid.uuid4().hex}.{ext}' if ext else uuid.uuid4().hex
    relative_path = f'{subfolder}/{unique_name}' if subfolder else unique_name

    content_type = (
        file_storage.mimetype
        or mimetypes.guess_type(original_name)[0]
        or 'application/octet-stream'
    )
    data = file_storage.read()
    file_storage.seek(0)
    return save_bytes(data, relative_path, content_type)


def delete_uploaded_file(relative_path):
    if not relative_path:
        return
    if storage_enabled():
        _storage_delete(relative_path)
        return
    full_path = os.path.join(current_app.config['UPLOAD_FOLDER'], relative_path)
    if os.path.exists(full_path):
        try:
            os.remove(full_path)
        except OSError:
            current_app.logger.warning('Impossible de supprimer %s', full_path)


def upload_url(relative_path):
    """URL a mettre dans un <img>/<a> pour un fichier uploade.

    - fichiers prives (requests/...) : route Flask protegee `private_file`
    - fichiers publics : URL publique Supabase (ou /static/uploads en local)
    """
    if not relative_path:
        return ''
    if _is_private(relative_path):
        return url_for('private_file', file_path=relative_path)
    if storage_enabled():
        cfg = current_app.config
        return (f"{cfg['SUPABASE_URL']}/storage/v1/object/public/"
                f"{cfg['SUPABASE_BUCKET']}/{relative_path}")
    return url_for('static', filename=f'uploads/{relative_path}')


def get_private_file_url(relative_path, expires_in=300):
    """URL temporaire (signee) vers un fichier prive.
    Retourne None si le fichier n'est pas disponible via Supabase."""
    if not storage_enabled():
        return None
    cfg = current_app.config
    url = f"{cfg['SUPABASE_URL']}/storage/v1/object/sign/{_bucket_for(relative_path)}/{relative_path}"
    resp = requests.post(
        url, json={'expiresIn': expires_in},
        headers=_storage_headers('application/json'), timeout=STORAGE_TIMEOUT)
    if resp.status_code >= 300:
        current_app.logger.error('URL signee refusee (%s) : %s', resp.status_code, resp.text)
        return None
    signed = resp.json().get('signedURL', '')
    if signed.startswith('/storage/v1'):
        return f"{cfg['SUPABASE_URL']}{signed}"
    return f"{cfg['SUPABASE_URL']}/storage/v1{signed}"


# ---------------------------------------------------------------------------
# Numeros de suivi & QR codes
# ---------------------------------------------------------------------------

def generate_tracking_number():
    """Genere un numero de suivi unique du type DEM-2026-A1B2C3."""
    year = datetime.utcnow().year
    suffix = ''.join(random.choices(string.ascii_uppercase + string.digits, k=6))
    return f'DEM-{year}-{suffix}'


def _qr_png_bytes(data):
    """Genere le QR code en memoire (aucune ecriture disque)."""
    buffer = io.BytesIO()
    qrcode.make(data).save(buffer, format='PNG')
    return buffer.getvalue()


def generate_qr_code(data, subfolder='qrcodes'):
    """Genere une image QR code encodant `data`, l'envoie dans le stockage
    et retourne son chemin relatif (ex. 'qrcodes/abc.png')."""
    relative_path = f'{subfolder}/{uuid.uuid4().hex}.png'
    return save_bytes(_qr_png_bytes(data), relative_path, 'image/png')


# ---------------------------------------------------------------------------
# Generation de documents PDF (ReportLab)
# ---------------------------------------------------------------------------

def generate_request_pdf(service_request):
    """Genere un recepisse PDF pour une demande administrative et retourne
    un buffer BytesIO pret a etre envoye via send_file()."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    # En-tete institutionnel
    c.setFillColor(colors.HexColor('#0B3D91'))
    c.rect(0, height - 30 * mm, width, 30 * mm, fill=True, stroke=False)
    c.setFillColor(colors.white)
    c.setFont('Helvetica-Bold', 16)
    c.drawString(20 * mm, height - 15 * mm, 'Plateforme de Digitalisation des Services Communaux')
    c.setFont('Helvetica', 10)
    c.drawString(20 * mm, height - 22 * mm, 'Récépissé de demande administrative')

    c.setFillColor(colors.black)
    y = height - 45 * mm
    c.setFont('Helvetica-Bold', 12)
    c.drawString(20 * mm, y, f'Numéro de suivi : {service_request.tracking_number}')
    y -= 8 * mm
    c.setFont('Helvetica', 11)
    c.drawString(20 * mm, y, f'Service : {service_request.service.name_fr}')
    y -= 7 * mm
    c.drawString(20 * mm, y, f'Demandeur : {service_request.citizen.user.full_name}')
    y -= 7 * mm
    c.drawString(20 * mm, y, f"Date de dépôt : {service_request.created_at.strftime('%d/%m/%Y %H:%M')}")
    y -= 7 * mm
    c.drawString(20 * mm, y, f'Statut actuel : {service_request.status_label()}')
    y -= 10 * mm

    if service_request.notes:
        c.setFont('Helvetica-Bold', 11)
        c.drawString(20 * mm, y, 'Description :')
        y -= 6 * mm
        c.setFont('Helvetica', 10)
        for line in _wrap_text(service_request.notes, 95):
            c.drawString(20 * mm, y, line)
            y -= 5 * mm

    # QR code (regenere en memoire : rien a lire sur le disque)
    if service_request.qr_code:
        qr_url = url_for('track_request',
                         tracking_number=service_request.tracking_number, _external=True)
        qr_image = ImageReader(io.BytesIO(_qr_png_bytes(qr_url)))
        c.drawImage(qr_image, width - 55 * mm, height - 85 * mm, 35 * mm, 35 * mm)

    c.setFont('Helvetica-Oblique', 8)
    c.drawString(20 * mm, 15 * mm,
                 "Ce document est généré automatiquement et ne nécessite pas de signature.")

    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer


def _wrap_text(text, max_chars):
    words = text.split()
    lines, current = [], ''
    for word in words:
        if len(current) + len(word) + 1 <= max_chars:
            current = f'{current} {word}'.strip()
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


# ---------------------------------------------------------------------------
# Export Excel (OpenPyXL)
# ---------------------------------------------------------------------------

def export_requests_to_excel(requests_list):
    """Cree un classeur Excel listant les demandes fournies et retourne un
    buffer BytesIO pret a etre envoye via send_file()."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Demandes'

    headers = ['N° de suivi', 'Citoyen', 'Service', 'Statut', 'Employé affecté', 'Date de dépôt']
    header_fill = PatternFill(start_color='0B3D91', end_color='0B3D91', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True)

    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font

    for row, req in enumerate(requests_list, start=2):
        ws.cell(row=row, column=1, value=req.tracking_number)
        ws.cell(row=row, column=2, value=req.citizen.user.full_name)
        ws.cell(row=row, column=3, value=req.service.name_fr)
        ws.cell(row=row, column=4, value=req.status_label())
        ws.cell(row=row, column=5, value=req.employee.user.full_name if req.employee else '-')
        ws.cell(row=row, column=6, value=req.created_at.strftime('%d/%m/%Y %H:%M'))

    for col_letter in ('A', 'B', 'C', 'D', 'E', 'F'):
        ws.column_dimensions[col_letter].width = 24

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


# ---------------------------------------------------------------------------
# Jetons securises (verification e-mail, reinitialisation mot de passe)
# ---------------------------------------------------------------------------

def get_serializer():
    return URLSafeTimedSerializer(current_app.config['SECRET_KEY'])


def generate_token(email, salt):
    return get_serializer().dumps(email, salt=salt)


def confirm_token(token, salt, max_age=86400):
    try:
        return get_serializer().loads(token, salt=salt, max_age=max_age)
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Envoi d'e-mails (Flask-Mail)
# ---------------------------------------------------------------------------

def send_email(subject, recipients, body, html=None):
    from app import mail  # import tardif pour eviter les imports circulaires
    if isinstance(recipients, str):
        recipients = [recipients]
    msg = Message(subject=subject, recipients=recipients, body=body, html=html)
    try:
        mail.send(msg)
    except Exception as exc:  # pragma: no cover - depend de la configuration SMTP
        current_app.logger.error('Échec de l\'envoi de l\'e-mail : %s', exc)


# ---------------------------------------------------------------------------
# Notifications internes
# ---------------------------------------------------------------------------

def create_notification(user, title, message='', link=None):
    from models import Notification, db
    notif = Notification(user_id=user.id, title=title, message=message, link=link)
    db.session.add(notif)
    db.session.commit()
    return notif


# ---------------------------------------------------------------------------
# Journalisation (piste d'audit)
# ---------------------------------------------------------------------------

def log_action(action, details=None, user=None):
    from models import Log, db
    actor = user or (current_user if current_user.is_authenticated else None)
    entry = Log(
        user_id=actor.id if actor else None,
        action=action,
        details=details,
        ip_address=request.remote_addr,
    )
    db.session.add(entry)
    db.session.commit()


# ---------------------------------------------------------------------------
# Decorateurs de controle d'acces par role
# ---------------------------------------------------------------------------

def roles_required(*role_names):
    def decorator(view_func):
        @wraps(view_func)
        def wrapped(*args, **kwargs):
            if not current_user.is_authenticated:
                return redirect(url_for('login', next=request.path))
            if not current_user.has_role(*role_names):
                abort(403)
            return view_func(*args, **kwargs)
        return wrapped
    return decorator


def admin_required(view_func):
    return roles_required('admin')(view_func)


def employee_required(view_func):
    return roles_required('admin', 'employee')(view_func)


def citizen_required(view_func):
    return roles_required('citizen')(view_func)


# ---------------------------------------------------------------------------
# Divers
# ---------------------------------------------------------------------------

def flash_errors(form):
    """Ajoute chaque erreur de validation d'un formulaire comme message flash."""
    for field, errors in form.errors.items():
        for error in errors:
            label = getattr(form, field).label.text
            flash(f'{label} : {error}', 'danger')
